import os
import json
import pypdf
import docx

manifest_path = r"e:\CIVCLENS\data\raw\manifestos\manifest.json"
raw_dir = r"e:\CIVCLENS\data\raw\manifestos\real_manifestos_20261006"

with open(manifest_path, "r", encoding="utf-8") as f:
    manifest = json.load(f)

print(f"Total files in manifest: {len(manifest['files'])}\n")

pdf_text_counts = []
for item in manifest["files"]:
    rel_path = item["relative_path"]
    full_path = os.path.join(raw_dir, rel_path)
    file_type = item["file_type"]

    if file_type == "pdf":
        try:
            reader = pypdf.PdfReader(full_path)
            total_pages = len(reader.pages)
            total_chars = 0
            non_empty_pages = 0
            for page in reader.pages:
                txt = page.extract_text() or ""
                if len(txt.strip()) > 20:
                    non_empty_pages += 1
                    total_chars += len(txt)
            pdf_text_counts.append({
                "filename": item["original_filename"],
                "pages": total_pages,
                "text_pages": non_empty_pages,
                "total_chars": total_chars,
                "ocr_needed": non_empty_pages < total_pages * 0.2
            })
        except Exception as e:
            pdf_text_counts.append({
                "filename": item["original_filename"],
                "pages": 0,
                "text_pages": 0,
                "total_chars": 0,
                "error": str(e),
                "ocr_needed": True
            })
    elif file_type == "docx":
        try:
            d = docx.Document(full_path)
            paragraphs = [p.text.strip() for p in d.paragraphs if p.text.strip()]
            total_chars = sum(len(p) for p in paragraphs)
            pdf_text_counts.append({
                "filename": item["original_filename"],
                "pages": len(d.paragraphs),
                "text_pages": len(paragraphs),
                "total_chars": total_chars,
                "ocr_needed": total_chars < 50
            })
        except Exception as e:
            pdf_text_counts.append({
                "filename": item["original_filename"],
                "pages": 0,
                "text_pages": 0,
                "total_chars": 0,
                "error": str(e),
                "ocr_needed": True
            })

for res in pdf_text_counts:
    print(f"File: {res['filename']}")
    print(f"  Pages/Paragraphs: {res['pages']} | Non-empty: {res['text_pages']} | Total Chars: {res['total_chars']:,} | OCR Needed: {res['ocr_needed']}\n")
