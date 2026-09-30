import docx

def inspect():
    try:
        doc = docx.Document(r'E:\CIVCLENS\CivicLens_TN_Project_Abstract.docx')
        print(f'Total tables found: {len(doc.tables)}')
        
        for i, table in enumerate(doc.tables):
            print(f'\n--- Table {i+1} ---')
            for row in table.rows:
                row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(' | '.join(row_text))
    except Exception as e:
        print(f'Error: {e}')

if __name__ == '__main__':
    inspect()
