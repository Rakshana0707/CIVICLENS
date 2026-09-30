import docx

def modify_doc():
    doc = docx.Document(r'E:\CIVCLENS\CivicLens_TN_Project_Abstract.docx')
    
    combined_data = [
        ("Module", "Core Function", "ML Algorithms & Techniques"),
        ("Budget & Scheme Analyzer", "Tracks department/scheme allocation over time; clusters related schemes; surfaces spending trends.", "Unsupervised Learning (K-Means, DBSCAN, Topic Modeling, PCA)"),
        ("Political Promise & Feasibility Analyzer", "Extracts claim structure, matches against historical promises, estimates cost, compares against budget.", "Supervised Learning (Logistic Regression, Random Forest, XGBoost, SVM); Deep Learning / NLP (IndicBERT, Sentence-BERT)"),
        ("Tamil News Framing Analyzer", "Compares sentiment, word choice, and emphasis across outlets covering the same event.", "Deep Learning / NLP (IndicBERT, Sentence-BERT); Unsupervised Learning (Topic Modeling)"),
        ("Representative Performance Dashboard", "Computes a configurable, weighted performance score from attendance, questions, debates, bills.", "Data Aggregation & Statistical Scoring"),
        ("Political Funding Transparency", "Visualizes party/donor funding and flags statistically unusual donation patterns.", "Anomaly Detection (Isolation Forest)"),
        ("Claim Verification Engine", "Links a political claim to budget evidence and historical data for an evidence-based assessment.", "Deep Learning / NLP (IndicBERT, Sentence-BERT)"),
        ("Recommendation (Optional)", "Personalizes which civic information is surfaced to a returning user.", "Reinforcement Learning (Multi-Armed Bandit, Q-Learning)")
    ]
    
    target_para = doc.paragraphs[2] 
    new_para = target_para.insert_paragraph_before('Combined Overview of Modules and ML Algorithms:')
    
    new_table = doc.add_table(rows=len(combined_data), cols=3)
    # Give table borders if possible
    new_table.style = doc.tables[0].style if len(doc.tables) > 1 else None
    
    for i, row_data in enumerate(combined_data):
        row = new_table.rows[i]
        for j, cell_text in enumerate(row_data):
            row.cells[j].text = cell_text
            
    p = new_para._p
    tbl = new_table._tbl
    p.addnext(tbl)
    
    doc.save(r'E:\CIVCLENS\CivicLens_TN_Project_Abstract.docx')
    print("Document saved successfully.")

if __name__ == '__main__':
    modify_doc()
