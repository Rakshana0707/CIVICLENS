import docx
from docx.shared import Pt, Inches

def create_submission():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('CivicLens TN - Project Abstract Submission', 0)
    title.alignment = 1 # Center
    
    # Combined Table
    doc.add_heading('1. Overview of Modules and ML Algorithms', level=1)
    
    combined_data = [
        ("Module", "Core Function", "ML Algorithms & Techniques"),
        ("Budget & Scheme Analyzer", "Tracks department/scheme allocation over time; clusters related schemes; surfaces spending trends.", "Unsupervised Learning (K-Means, DBSCAN, Topic Modeling, PCA)"),
        ("Political Promise & Feasibility Analyzer", "Extracts claim structure, matches against historical promises, estimates cost, compares against budget.", "Supervised Learning (Logistic Regression, Random Forest, XGBoost, SVM); Deep Learning / NLP (IndicBERT, Sentence-BERT)"),
        ("Tamil News Framing Analyzer", "Compares sentiment, word choice, and emphasis across outlets covering the same event.", "Deep Learning / NLP (IndicBERT, Sentence-BERT); Unsupervised Learning (Topic Modeling)"),
        ("Representative Performance Dashboard", "Computes a configurable, weighted performance score from attendance, questions, debates, bills.", "Data Aggregation & Statistical Scoring"),
        ("Political Funding Transparency", "Visualizes party/donor funding and flags statistically unusual donation patterns.", "Anomaly Detection (Isolation Forest)"),
        ("Claim Verification Engine", "Links a political claim to budget evidence and historical data for an evidence-based assessment.", "Deep Learning / NLP (IndicBERT, Sentence-BERT)")
    ]
    
    table = doc.add_table(rows=len(combined_data), cols=3)
    table.style = 'Table Grid'
    
    for i, row_data in enumerate(combined_data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            row.cells[j].text = cell_text
            
    doc.add_paragraph() # spacing
            
    # Objectives
    doc.add_heading('2. Objectives', level=1)
    objectives = [
        "Collect and structure Tamil Nadu department- and scheme-wise budget data across multiple years.",
        "Cluster and track schemes over time using NLP and unsupervised learning to surface spending trends.",
        "Extract structured claims (target group, benefit, amount, time period, sector) from political statements using NLP.",
        "Compare new promises and claims against historical schemes and actual budget data to generate a feasibility/evidence assessment.",
        "Analyze Tamil news articles covering the same event to measure sentiment, framing, and emphasis across sources.",
        "Build a configurable, weighted performance score for MLAs/MPs from attendance, questions, debates, bills, and constituency indicators.",
        "Apply anomaly detection to publicly available political-funding data to flag statistically unusual donation patterns.",
        "Present all outputs through a unified citizen dashboard with explainable, evidence-linked results."
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')
        
    # ML/DL Methods
    doc.add_heading('3. ML/DL Methods and Algorithms', level=1)
    methods = [
        "Supervised Learning: Logistic Regression, Random Forest, XGBoost, SVM (used for claim/promise implementation-status classification).",
        "Unsupervised Learning: K-Means, DBSCAN, Topic Modeling, PCA (used for scheme clustering and news-topic/framing discovery).",
        "Deep Learning / NLP: IndicBERT / multilingual BERT, Sentence-BERT embeddings (used for Tamil sentiment analysis, semantic similarity for promise matching, claim/entity extraction).",
        "Anomaly Detection: Isolation Forest (used for flagging statistically unusual political-donation patterns).",
        "Reinforcement Learning: Multi-Armed Bandit / Q-Learning (optional module for personalizing surfaced civic information)."
    ]
    for method in methods:
        doc.add_paragraph(method, style='List Bullet')
        
    # Datasets
    doc.add_heading('4. Corresponding Datasets', level=1)
    doc.add_paragraph("The project prioritizes real, publicly available data. Where domain-specific labeled data is not publicly available, a smaller annotated subset will be manually created for model training.")
    datasets = [
        "Tamil Nadu Budget Portal (Public): Department-wise demands, scheme-level allocations, and expenditure by year.",
        "PRS Legislative Research (Public): MLA/MP attendance, questions asked, debate participation, and bill activity.",
        "ADR / MyNeta (Public): Candidate financial disclosures and political-party donation/funding data.",
        "Tamil News Archives (Public/Scraped): Articles from multiple Tamil-language outlets covering the same political events for framing/sentiment comparison.",
        "Political Manifestos & Past Scheme Documents (Public): Used to match new promises against historical precedents."
    ]
    for ds in datasets:
        doc.add_paragraph(ds, style='List Bullet')
        
    # References
    doc.add_heading('5. References', level=1)
    refs = [
        "Government of Tamil Nadu. (n.d.). Tamil Nadu State Budget Portal. Retrieved from https://www.tn.gov.in/",
        "PRS Legislative Research. (n.d.). State Legislative Data - Tamil Nadu. Retrieved from https://prsindia.org/",
        "Association for Democratic Reforms (ADR). (n.d.). MyNeta Database. Retrieved from https://myneta.info/",
        "Devlin, J. et al. (2018). BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding. arXiv:1810.04805.",
        "Kakwani, D. et al. (2020). IndicNLPSuite: Monolingual Corpora, Evaluation Benchmarks and Pre-trained Multilingual Language Models for Indian Languages. Findings of EMNLP.",
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830."
    ]
    for i, ref in enumerate(refs):
        doc.add_paragraph(f"[{i+1}] {ref}")
        
    doc.save(r'E:\CIVCLENS\CivicLens_TN_Abstract_Submission.docx')
    print("Document successfully created!")

if __name__ == '__main__':
    create_submission()
